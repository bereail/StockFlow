import csv
import io
import logging

from django.core.paginator import Paginator
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Count, Prefetch
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.timezone import is_naive, make_aware
from django.views.decorators.http import require_POST
from ..models import Impresora, AsignacionImpresora, Movimiento, MovimientoDetalle, Reparacion
from ..forms.impresoras import ImpresoraForm, EntregaRapidaImpresoraForm
from ..forms.asignaciones import AsignacionImpresoraForm
from ..services.items import item_de_impresora
from ..services.impresoras import asignar_impresora_a_servicio
from ..services.busqueda import buscar_texto
from ..services.listados import ordenar
from ..services.red import hacer_ping

logger = logging.getLogger(__name__)


@login_required
def impresoras_page(request):
    q = (request.GET.get("q") or "").strip()
    estado = (request.GET.get("estado") or "").strip()

    impresoras = (
        Impresora.objects
        .select_related("toner", "articulo")
        .prefetch_related(
            Prefetch(
                "asignaciones",
                queryset=AsignacionImpresora.objects.select_related("servicio").order_by("-fecha_desde"),
            )
        )
    )

    impresoras = buscar_texto(impresoras, q, "marca", "modelo", "patrimonio", "ip")
    if estado:
        impresoras = impresoras.filter(estado=estado)

    impresoras, sort_actual, dir_actual = ordenar(
        request, impresoras,
        campos={"marca": "marca", "modelo": "modelo", "ip": "ip", "patrimonio": "patrimonio", "estado": "estado"},
        default="marca",
    )

    paginator = Paginator(impresoras, 5)
    page_obj  = paginator.get_page(request.GET.get("page", 1))

    # IPs repetidas entre impresoras: alerta visual (no bloquea nada,
    # hay datos reales preexistentes con IP duplicada — ver memoria del proyecto).
    ips_duplicadas = set(
        Impresora.objects
        .exclude(ip__isnull=True).exclude(ip="")
        .values("ip")
        .annotate(n=Count("id"))
        .filter(n__gt=1)
        .values_list("ip", flat=True)
    )

    return render(request, "inventario/impresoras/impresoras.html", {
        "impresoras": page_obj, "page_obj": page_obj, "q": q,
        "estado": estado, "sort_actual": sort_actual, "dir_actual": dir_actual,
        "ips_duplicadas": ips_duplicadas,
    })


@login_required
def impresoras_exportar_excel(request):
    """Excel (.xlsx) con todas las impresoras: una fila por impresora con
    "Impresora <servicio>", IP y MAC — mismo formato que la planilla que ya
    se usa para repartir estos datos, ordenado por servicio."""
    from openpyxl import Workbook

    impresoras = (
        Impresora.objects
        .prefetch_related(
            Prefetch(
                "asignaciones",
                queryset=AsignacionImpresora.objects.select_related("servicio").order_by("-fecha_desde"),
            )
        )
    )

    filas = []
    for i in impresoras:
        servicio = i.servicio_actual
        etiqueta = f"Impresora {servicio.nombre}" if servicio else "Impresora (sin servicio)"
        filas.append((etiqueta, i.ip or "", i.mac or ""))
    filas.sort(key=lambda f: f[0])

    wb = Workbook()
    ws = wb.active
    ws.title = "Impresoras"
    for fila, (etiqueta, ip, mac) in enumerate(filas, start=1):
        ws.cell(row=fila, column=1, value=etiqueta)
        ws.cell(row=fila, column=2, value=ip)
        ws.cell(row=fila, column=3, value=mac)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 20

    nombre = f"impresoras_{timezone.now():%Y%m%d}.xlsx"
    buf = io.BytesIO()
    wb.save(buf)

    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{nombre}"'
    return response


@login_required
def impresora_detail(request, pk):
    impresora = get_object_or_404(Impresora.objects.select_related("toner", "articulo"), pk=pk)

    asignaciones = (
        impresora.asignaciones
        .select_related("servicio")
        .order_by("-fecha_desde")
    )

    item = item_de_impresora(impresora)
    reparaciones = (
        Reparacion.objects
        .filter(item=item)
        .select_related("proveedor")
        .order_by("-creado")
    )

    pcs_asociadas = impresora.computadoras.select_related("servicio").order_by("nombre_pc")

    return render(request, "inventario/impresoras/impresora_detail.html", {
        "impresora": impresora,
        "asignaciones": asignaciones,
        "reparaciones": reparaciones,
        "pcs_asociadas": pcs_asociadas,
    })


@login_required
def impresora_ping(request, pk):
    impresora = get_object_or_404(Impresora, pk=pk)
    if not impresora.ip:
        return JsonResponse({"online": False, "error": "Esta impresora no tiene IP cargada."}, status=400)
    return JsonResponse({"online": hacer_ping(impresora.ip)})


@login_required
def asignar_impresora(request, impresora_id):
    impresora = get_object_or_404(Impresora, id=impresora_id)

    if request.method == "POST":
        form = AsignacionImpresoraForm(request.POST)
        if form.is_valid():
            hoy = timezone.localdate()

            AsignacionImpresora.objects.filter(
                impresora=impresora,
                fecha_hasta__isnull=True
            ).update(fecha_hasta=hoy)

            asignacion = form.save(commit=False)
            asignacion.impresora = impresora
            if not asignacion.fecha_desde:
                asignacion.fecha_desde = hoy
            asignacion.save()

            messages.success(request, f"Impresora asignada a {asignacion.servicio} correctamente.")
            return redirect("impresoras_page")
    else:
        form = AsignacionImpresoraForm(initial={"fecha_desde": timezone.localdate()})

    return render(
        request,
        "inventario/impresoras/asignar_impresora.html",
        {"impresora": impresora, "form": form, "title": "Asignar impresora a servicio"},
    )


@login_required
def impresora_create(request):
    if request.method == "POST":
        form = ImpresoraForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Impresora creada.")
            return redirect("impresoras_page")
        messages.error(request, "Revisá los errores del formulario.")
    else:
        form = ImpresoraForm()

    return render(request, "inventario/impresoras/impresora_form.html", {"form": form, "mode": "create"})


@login_required
def impresora_edit(request, pk):
    impresora = get_object_or_404(Impresora, pk=pk)

    if request.method == "POST":
        form = ImpresoraForm(request.POST, instance=impresora)
        if form.is_valid():
            form.save()
            messages.success(request, "Impresora actualizada.")
            return redirect("impresoras_page")
        messages.error(request, "Revisá los errores del formulario.")
    else:
        form = ImpresoraForm(instance=impresora)

    return render(request, "inventario/impresoras/impresora_form.html", {"form": form, "mode": "edit", "impresora": impresora})


@login_required
@require_POST
def impresora_toggle(request, pk):
    impresora = get_object_or_404(Impresora, pk=pk)
    impresora.estado = "INACTIVA" if impresora.estado == "ACTIVA" else "ACTIVA"
    impresora.save(update_fields=["estado"])
    return redirect("impresoras_page")


@login_required
def impresora_entrega(request):
    if request.method == "POST":
        form = EntregaRapidaImpresoraForm(request.POST)
        if form.is_valid():
            servicio = form.cleaned_data["servicio"]
            impresora = form.cleaned_data["impresora"]
            observaciones = (form.cleaned_data.get("observaciones") or "").strip()
            fecha = form.cleaned_data.get("fecha") or timezone.now()

            if is_naive(fecha):
                fecha = make_aware(fecha, timezone.get_current_timezone())

            try:
                with transaction.atomic():
                    movimiento = Movimiento.objects.create(
                        tipo="EGRESO",
                        fecha=fecha,
                        servicio=servicio,
                        observaciones=observaciones,
                    )

                    MovimientoDetalle.objects.create(
                        movimiento=movimiento,
                        item=item_de_impresora(impresora),
                        cantidad=1,
                    )

                    # La entrega es, en la práctica, una asignación a ese
                    # servicio: se refleja igual que si se hiciera desde
                    # "Asignar / Mover" o desde la carga de Patrimonio.
                    asignar_impresora_a_servicio(
                        impresora,
                        servicio,
                        fecha=fecha.date(),
                        observaciones=observaciones,
                    )

                messages.success(request, "Movimiento de impresora registrado correctamente.")
                return redirect("impresoras_page")

            except Exception:
                logger.exception("Error al registrar movimiento de entrega de impresora")
                messages.error(request, "No se pudo registrar el movimiento. Intentá de nuevo.")
        else:
            messages.error(request, "Revisá los datos del formulario.")
    else:
        form = EntregaRapidaImpresoraForm()

    return render(request, "inventario/impresoras/impresora_entrega.html", {"form": form})


@login_required
def impresora_historial(request):
    q = (request.GET.get("q") or "").strip()
    impresora_id = (request.GET.get("impresora") or "").strip()

    detalles = (
        MovimientoDetalle.objects
        .select_related("movimiento", "item", "item__impresora", "movimiento__servicio")
        .filter(movimiento__tipo="EGRESO", item__tipo="IMPRESORA")
        .order_by("-movimiento__fecha")
    )

    detalles = buscar_texto(
        detalles, q,
        "item__impresora__marca", "item__impresora__modelo",
        "movimiento__servicio__nombre", "movimiento__observaciones",
    )
    if impresora_id:
        detalles = detalles.filter(item__impresora_id=impresora_id)

    paginator = Paginator(detalles, 5)
    page_obj  = paginator.get_page(request.GET.get("page", 1))
    return render(request, "inventario/impresoras/impresoras_historial.html", {
        "detalles": page_obj, "page_obj": page_obj, "q": q, "impresora_id": impresora_id,
    })
